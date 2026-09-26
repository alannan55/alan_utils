# -*- coding: utf-8 -*-
"""
按供应商拆分Excel表格工具
支持根据供应商名称拆分表格，并支持简写映射和特殊规则
"""

import json
import os
import sys
import traceback
from copy import copy
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

CONFIG_FILE_NAME = "supplier_shortnames.json"


DEFAULT_SUPPLIER_MAPPING = {
    "安华仁": "安华仁",
    "奥瑞康": "奥瑞康",
    "超伦贸易": "超伦方盛",
    "晨益药业": "晨益",
    "方盛康源": "超伦方盛",
    "惠生医药": "惠生",
    "嘉迪医疗": "嘉迪",
    "凯宏鑫": "凯宏鑫",
    "仟草": "仟草",
    "瑞康志德": "瑞康志德",
    "友孚": "友孚",
    "华润润采": "华润润采",
    "药九九": "药九九",
}
DEFAULT_KEEP_PRICE_KEYWORDS = ["晨益药业"]


# JSON配置文件路径（与脚本同级目录，或exe同级目录）
def get_external_config_file_path():
    """获取外部JSON配置文件路径，exe优先读取exe所在目录。"""
    if getattr(sys, 'frozen', False):
        return os.path.join(os.path.dirname(sys.executable), CONFIG_FILE_NAME)
    return os.path.join(os.path.dirname(__file__), CONFIG_FILE_NAME)


def get_bundled_config_file_path():
    """获取打包进exe内部的JSON配置文件路径。"""
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, CONFIG_FILE_NAME)
    return None


EXTERNAL_CONFIG_FILE = get_external_config_file_path()
BUNDLED_CONFIG_FILE = get_bundled_config_file_path()

SPECIAL_SUPPLIER_OUTPUTS = {
    "北京惠生医药有限责任公司": "惠生",
    "北京药九九医药科技有限公司": "药九九",
}
MERGED_SUPPLIER_OUTPUT_NAME = "惠生+药九九"


class SupplierSplitterGUI(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("按供应商拆分Excel表格")
        self.geometry("500x120")
        self.resizable(False, False)

        self.input_file = tk.StringVar()
        self.output_dir = tk.StringVar()
        self.supplier_mapping = {}  # 供应商名称到简写的映射
        self.keep_price_keywords = []  # 需要保留售价列的关键字列表

        # 加载JSON配置
        self.load_config()

        self._build_ui()

    def load_config(self):
        """优先加载外部JSON配置，失败时使用打包内置配置。"""
        self.supplier_mapping = DEFAULT_SUPPLIER_MAPPING.copy()
        self.keep_price_keywords = DEFAULT_KEEP_PRICE_KEYWORDS.copy()

        try:
            if os.path.exists(EXTERNAL_CONFIG_FILE):
                self.apply_config_from_file(EXTERNAL_CONFIG_FILE)
                return
        except Exception as e:
            messagebox.showwarning("警告", f"外部配置文件加载失败：{e}\n将使用内置配置。")

        if BUNDLED_CONFIG_FILE and os.path.exists(BUNDLED_CONFIG_FILE):
            try:
                self.apply_config_from_file(BUNDLED_CONFIG_FILE)
            except Exception:
                # 已经提前加载了代码内置默认配置，这里静默兜底即可。
                pass

    def apply_config_from_file(self, config_file):
        """从JSON配置文件加载供应商名称映射。"""
        with open(config_file, 'r', encoding='utf-8') as f:
            config = json.load(f)
            # 支持新格式（有mappings和keep_price）和旧格式（直接是映射字典）
            if isinstance(config, dict) and "mappings" in config:
                supplier_mapping = config.get("mappings", {})
                keep_price_keywords = config.get("keep_price", [])
                if not isinstance(supplier_mapping, dict):
                    raise ValueError("mappings 必须是JSON对象")
                if not isinstance(keep_price_keywords, list):
                    raise ValueError("keep_price 必须是JSON数组")
                self.supplier_mapping = supplier_mapping
                self.keep_price_keywords = keep_price_keywords
            elif isinstance(config, dict):
                # 兼容旧格式
                self.supplier_mapping = config
                self.keep_price_keywords = []
            else:
                raise ValueError("配置文件必须是JSON对象")

    def _build_ui(self):
        pad = 6

        frame_main = ttk.Frame(self, padding=pad)
        frame_main.pack(fill="both", expand=True)

        # 输入文件
        ttk.Label(frame_main, text="输入表格文件:").grid(column=0, row=0, sticky="w", pady=2)
        ttk.Entry(frame_main, textvariable=self.input_file, width=42).grid(column=1, row=0, padx=4, pady=2)
        ttk.Button(frame_main, text="选择文件", command=self.choose_input_file).grid(column=2, row=0, pady=2)

        # 输出目录
        ttk.Label(frame_main, text="保存路径:").grid(column=0, row=1, sticky="w", pady=2)
        ttk.Entry(frame_main, textvariable=self.output_dir, width=42).grid(column=1, row=1, padx=4, pady=2)
        ttk.Button(frame_main, text="选择目录", command=self.choose_output_dir).grid(column=2, row=1, pady=2)

        # 执行按钮
        frame_buttons = ttk.Frame(frame_main)
        frame_buttons.grid(column=0, columnspan=3, row=2, pady=8)
        ttk.Button(frame_buttons, text="开始拆分", command=self.split_files, width=17).pack(side="left", padx=4)
        ttk.Button(frame_buttons, text="退出", command=self.quit, width=17).pack(side="left", padx=4)

        # 状态栏
        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(frame_main, textvariable=self.status_var, foreground="gray").grid(column=0, columnspan=3, row=3, pady=3)

    def choose_input_file(self):
        """选择输入文件"""
        filename = filedialog.askopenfilename(
            title="选择Excel表格文件",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if filename:
            self.input_file.set(filename)

    def choose_output_dir(self):
        """选择输出目录"""
        dirname = filedialog.askdirectory(title="选择保存目录")
        if dirname:
            self.output_dir.set(dirname)

    def get_supplier_shortname(self, supplier_name):
        """根据供应商名称获取简写名称"""
        if not supplier_name:
            return "未知供应商"
        
        supplier_name_str = str(supplier_name).strip()

        if supplier_name_str in SPECIAL_SUPPLIER_OUTPUTS:
            return SPECIAL_SUPPLIER_OUTPUTS[supplier_name_str]
        
        # 查找完全匹配
        if supplier_name_str in self.supplier_mapping:
            return self.supplier_mapping[supplier_name_str]
        
        # 查找部分匹配：只要供应商名称包含JSON中的键，就返回对应的值
        # 例如：JSON中"奥瑞康":"奥瑞康"，则"北京奥瑞康科技发展有限公司"和"北京奥瑞康科技发展有限公司-月结"都会匹配为"奥瑞康"
        for key_name, short_name in self.supplier_mapping.items():
            if key_name in supplier_name_str:
                return short_name
        
        # 如果没有匹配到，返回原名称
        return supplier_name_str

    def should_keep_price_col(self, supplier_name):
        """判断是否需要保留售价列"""
        if not supplier_name:
            return False
        
        supplier_name_str = str(supplier_name).strip()
        
        # 检查供应商名称是否包含需要保留售价的关键字
        for keyword in self.keep_price_keywords:
            if keyword in supplier_name_str:
                return True
        
        return False

    def get_merged_supplier_sort_key(self, row_idx, sheet, serial_col_idx):
        """合并惠生和药九九时，按序号中的门店关键词排序。"""
        serial_value = sheet.cell(row=row_idx, column=serial_col_idx).value
        serial_value_str = str(serial_value or "")
        if "金顶街" in serial_value_str:
            return (0, row_idx)
        if "老山" in serial_value_str:
            return (2, row_idx)
        return (1, row_idx)

    def copy_cell_with_style_and_comment(self, source_cell, target_cell):
        """复制单元格内容、样式、超链接和备注，不重新设置任何颜色。"""
        target_cell.value = source_cell.value
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)
        if source_cell.hyperlink:
            target_cell._hyperlink = copy(source_cell.hyperlink)
        if source_cell.comment:
            target_cell.comment = copy(source_cell.comment)

    def copy_row_dimension(self, source_sheet, target_sheet, source_row_idx, target_row_idx):
        """复制行高、隐藏状态等行属性。"""
        source_dim = source_sheet.row_dimensions[source_row_idx]
        target_dim = target_sheet.row_dimensions[target_row_idx]
        target_dim.height = source_dim.height
        target_dim.hidden = source_dim.hidden
        target_dim.outlineLevel = source_dim.outlineLevel
        target_dim.collapsed = source_dim.collapsed

    def copy_column_dimension(self, source_sheet, target_sheet, source_col_idx, target_col_idx):
        """复制列宽、隐藏状态等列属性。"""
        source_letter = get_column_letter(source_col_idx)
        target_letter = get_column_letter(target_col_idx)
        source_dim = source_sheet.column_dimensions[source_letter]
        target_dim = target_sheet.column_dimensions[target_letter]
        target_dim.width = source_dim.width
        target_dim.hidden = source_dim.hidden
        target_dim.outlineLevel = source_dim.outlineLevel
        target_dim.collapsed = source_dim.collapsed

    def auto_adjust_column_width(self, sheet, num_columns):
        """自动调整列宽以完整显示内容"""
        for col_idx in range(1, num_columns + 1):
            max_length = 0
            # 检查表头和所有数据行
            for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=col_idx, max_col=col_idx):
                cell = row[0]
                if cell.value:
                    # 计算单元格内容的显示长度
                    cell_value = str(cell.value)
                    # 估算显示宽度：中文字符按2个单位，英文字符按1个单位
                    display_length = 0
                    for char in cell_value:
                        # 判断是否为中文字符（包括中文标点等宽字符）
                        if ord(char) > 127:
                            display_length += 2  # 中文字符通常需要2个字符宽度
                        else:
                            display_length += 1  # 英文字符需要1个字符宽度
                    max_length = max(max_length, display_length)
            
            # 设置列宽（openpyxl的列宽单位近似等于字符数）
            # 加上一些边距（+3），最小宽度为10
            adjusted_width = max(max_length + 3, 10)
            # 获取列字母并设置列宽
            col_letter = get_column_letter(col_idx)
            sheet.column_dimensions[col_letter].width = adjusted_width

    def split_files(self):
        """执行拆分操作"""
        input_path = self.input_file.get()
        output_path = self.output_dir.get()

        if not input_path or not os.path.exists(input_path):
            messagebox.showerror("错误", "请选择有效的输入文件！")
            return

        if not output_path:
            messagebox.showerror("错误", "请选择保存路径！")
            return

        try:
            self.status_var.set("正在处理...")
            self.update()

            # 读取Excel文件
            workbook = load_workbook(filename=input_path)
            sheet = workbook.active

            # 读取第一行，找到"供应商名称"列的索引
            header_row = []
            supplier_col_idx = None
            serial_col_idx = None
            
            for col_idx, cell in enumerate(sheet[1], start=1):
                header_value = str(cell.value) if cell.value else ""
                header_row.append(cell.value if cell.value else "")
                if "供应商名称" in header_value:
                    supplier_col_idx = col_idx
                if "序号" in header_value:
                    serial_col_idx = col_idx

            if supplier_col_idx is None:
                messagebox.showerror("错误", "未找到包含'供应商名称'的列！")
                self.status_var.set("错误：未找到供应商名称列")
                return

            # 找到"售价"列的索引（用于晨益药业的特殊处理）
            price_col_idx = None
            for col_idx, cell in enumerate(sheet[1], start=1):
                if cell.value and "售价" in str(cell.value):
                    price_col_idx = col_idx
                    break

            # 读取所有数据行
            data_rows = []
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), start=2):
                supplier_cell = sheet.cell(row=row_idx, column=supplier_col_idx)
                supplier_name = supplier_cell.value if supplier_cell.value else None
                if supplier_name:
                    data_rows.append((row_idx, supplier_name))

            input_supplier_names = {str(supplier_name).strip() for _, supplier_name in data_rows}
            merge_special_suppliers = set(SPECIAL_SUPPLIER_OUTPUTS).issubset(input_supplier_names)

            # 按供应商简写名称分组（相同简写的供应商会合并到同一个文件）
            supplier_groups = {}
            for row_idx, supplier_name in data_rows:
                supplier_name_str = str(supplier_name).strip()
                if merge_special_suppliers and supplier_name_str in SPECIAL_SUPPLIER_OUTPUTS:
                    short_name = MERGED_SUPPLIER_OUTPUT_NAME
                else:
                    short_name = self.get_supplier_shortname(supplier_name_str)
                
                # 使用简写名称作为分组键
                if short_name not in supplier_groups:
                    supplier_groups[short_name] = []
                supplier_groups[short_name].append(row_idx)

            if not supplier_groups:
                messagebox.showwarning("警告", "未找到任何供应商数据！")
                self.status_var.set("警告：未找到供应商数据")
                return

            # 检查哪些文件会被覆盖
            files_to_overwrite = []
            for short_name in supplier_groups.keys():
                output_file = os.path.join(output_path, f"{short_name}.xlsx")
                if os.path.exists(output_file):
                    files_to_overwrite.append(f"{short_name}.xlsx")

            # 如果有文件会被覆盖，显示警告并询问用户
            if files_to_overwrite:
                overwrite_message = f"以下文件已存在，将被覆盖：\n\n" + "\n".join(files_to_overwrite)
                if not messagebox.askyesno("确认覆盖", overwrite_message + "\n\n是否继续？"):
                    self.status_var.set("操作已取消")
                    return

            # 为每个供应商创建新的Excel文件
            saved_files = []
            for short_name, row_indices in supplier_groups.items():
                if short_name == MERGED_SUPPLIER_OUTPUT_NAME and serial_col_idx:
                    row_indices = sorted(
                        row_indices,
                        key=lambda row_idx: self.get_merged_supplier_sort_key(row_idx, sheet, serial_col_idx)
                    )

                # 创建新工作簿
                new_wb = Workbook()
                new_sheet = new_wb.active

                # 判断是否需要保留"售价"列（通过检查第一个供应商名称来判断）
                first_supplier_name = sheet.cell(row=row_indices[0], column=supplier_col_idx).value
                keep_price_col = self.should_keep_price_col(first_supplier_name)

                # 写入表头
                header_to_write = []
                header_col_mapping = {}  # 原列索引 -> 新列索引的映射
                new_col_idx = 1

                for col_idx, header_value in enumerate(header_row, start=1):
                    # 如果是售价列且不需要保留，则跳过
                    if price_col_idx and col_idx == price_col_idx and not keep_price_col:
                        continue
                    
                    header_to_write.append(header_value)
                    header_col_mapping[col_idx] = new_col_idx
                    new_col_idx += 1

                # 写入表头到新表格
                for original_col_idx, target_col_idx in header_col_mapping.items():
                    self.copy_cell_with_style_and_comment(
                        sheet.cell(row=1, column=original_col_idx),
                        new_sheet.cell(row=1, column=target_col_idx)
                    )
                    self.copy_column_dimension(sheet, new_sheet, original_col_idx, target_col_idx)
                self.copy_row_dimension(sheet, new_sheet, 1, 1)

                # 写入数据行
                for new_row_idx, original_row_idx in enumerate(row_indices, start=2):
                    new_col_idx = 1
                    for col_idx, header_value in enumerate(header_row, start=1):
                        # 如果是售价列且不需要保留，则跳过
                        if price_col_idx and col_idx == price_col_idx and not keep_price_col:
                            continue
                        
                        self.copy_cell_with_style_and_comment(
                            sheet.cell(row=original_row_idx, column=col_idx),
                            new_sheet.cell(row=new_row_idx, column=new_col_idx)
                        )
                        new_col_idx += 1

                    self.copy_row_dimension(sheet, new_sheet, original_row_idx, new_row_idx)

                # 自动调整列宽
                self.auto_adjust_column_width(new_sheet, len(header_to_write))
                
                # 保存文件（直接覆盖，不重命名）
                output_file = os.path.join(output_path, f"{short_name}.xlsx")
                new_wb.save(output_file)
                saved_files.append(output_file)

            # 显示完成消息
            messagebox.showinfo("完成", f"拆分完成！\n共生成 {len(saved_files)} 个文件：\n" + 
                              "\n".join([os.path.basename(f) for f in saved_files]))
            self.status_var.set(f"完成：已生成 {len(saved_files)} 个文件")

        except Exception as e:
            traceback.print_exc()
            messagebox.showerror("错误", f"处理过程中发生错误：\n{e}")
            self.status_var.set(f"错误：{str(e)}")


def main():
    app = SupplierSplitterGUI()
    app.mainloop()


if __name__ == "__main__":
    main()
