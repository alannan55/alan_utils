# -*- coding: utf-8 -*-
"""
比较两个Excel表格并标记不匹配的条目
A表（进货表）和B表（供应商出货表）的字段对应关系：
- A表"进价" ↔ B表"销售价"
- A表"进货数量" ↔ B表"数量"
- A表"进货金额" ↔ B表"销售金额"
"""

import sys
import os
import traceback
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from typing import List, Tuple, Optional
import tkinter as tk
from tkinter import ttk, filedialog, messagebox


def find_column_index(sheet, header_text: str) -> Optional[int]:
    """在表头行中查找包含指定文本的列索引（从1开始）"""
    for col_idx, cell in enumerate(sheet[1], start=1):
        if cell.value and header_text in str(cell.value):
            return col_idx
    return None


def identify_table_type(sheet) -> Optional[str]:
    """
    自动识别表格类型
    返回 'A' 如果包含"进价"、"进货数量"、"进货金额"
    返回 'B' 如果包含"销售价"、"数量"、"销售金额"
    返回 None 如果都不匹配
    """
    has_price_a = find_column_index(sheet, "进价") is not None
    has_qty_a = find_column_index(sheet, "进货数量") is not None
    has_amount_a = find_column_index(sheet, "进货金额") is not None
    
    has_price_b = find_column_index(sheet, "销售价") is not None
    has_qty_b = find_column_index(sheet, "数量") is not None
    has_amount_b = find_column_index(sheet, "销售金额") is not None
    
    # 检查是否为A表（进货表）
    if has_price_a and has_qty_a and has_amount_a:
        return 'A'
    
    # 检查是否为B表（销售表）
    if has_price_b and has_qty_b and has_amount_b:
        return 'B'
    
    return None


def get_cell_value(sheet, row: int, col: int):
    """获取单元格的值，如果为空则返回None"""
    cell = sheet.cell(row=row, column=col)
    return cell.value if cell.value is not None else None


def compare_and_mark(path_file1: str, path_file2: str, save_path: str, status_callback=None):
    """
    比较两个Excel表格，标记不匹配的条目
    自动识别哪个是A表（进货表），哪个是B表（销售表）
    
    参数:
        path_file1: 第一个表格文件路径
        path_file2: 第二个表格文件路径
        save_path: 保存路径（目录）
        status_callback: 可选的状态回调函数，用于GUI更新状态
    """
    def log(message):
        """日志输出"""
        print(message)
        if status_callback:
            status_callback(message)
    
    # 检查文件是否存在
    if not os.path.exists(path_file1):
        error_msg = f"错误：文件不存在：{path_file1}"
        log(error_msg)
        if status_callback is None:
            raise FileNotFoundError(error_msg)
        return False
    if not os.path.exists(path_file2):
        error_msg = f"错误：文件不存在：{path_file2}"
        log(error_msg)
        if status_callback is None:
            raise FileNotFoundError(error_msg)
        return False
    if not os.path.exists(save_path):
        error_msg = f"错误：保存路径不存在：{save_path}"
        log(error_msg)
        if status_callback is None:
            raise FileNotFoundError(error_msg)
        return False
    
    # 加载工作簿
    log(f"正在加载文件1：{path_file1}")
    wb1 = load_workbook(filename=path_file1, read_only=False, data_only=True)
    ws1 = wb1.active
    
    log(f"正在加载文件2：{path_file2}")
    wb2 = load_workbook(filename=path_file2, read_only=False, data_only=True)
    ws2 = wb2.active
    
    # 自动识别表格类型
    type1 = identify_table_type(ws1)
    type2 = identify_table_type(ws2)
    
    if type1 is None:
        error_msg = f"错误：无法识别文件1的表格类型（应包含'进价'、'进货数量'、'进货金额'或'销售价'、'数量'、'销售金额'）"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    
    if type2 is None:
        error_msg = f"错误：无法识别文件2的表格类型（应包含'进价'、'进货数量'、'进货金额'或'销售价'、'数量'、'销售金额'）"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    
    if type1 == type2:
        error_msg = f"错误：两个表格类型相同（都是{'进货表' if type1 == 'A' else '销售表'}），无法进行比较"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    
    # 确定哪个是A表，哪个是B表
    if type1 == 'A':
        wb_a, ws_a = wb1, ws1
        wb_b, ws_b = wb2, ws2
        path_a, path_b = path_file1, path_file2
        log(f"已识别：文件1为A表（进货表），文件2为B表（销售表）")
    else:
        wb_a, ws_a = wb2, ws2
        wb_b, ws_b = wb1, ws1
        path_a, path_b = path_file2, path_file1
        log(f"已识别：文件2为A表（进货表），文件1为B表（销售表）")
    
    # 查找A表的列索引
    price_col_a = find_column_index(ws_a, "进价")
    qty_col_a = find_column_index(ws_a, "进货数量")
    amount_col_a = find_column_index(ws_a, "进货金额")
    
    if not price_col_a:
        error_msg = "错误：在A表中未找到'进价'列"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    if not qty_col_a:
        error_msg = "错误：在A表中未找到'进货数量'列"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    if not amount_col_a:
        error_msg = "错误：在A表中未找到'进货金额'列"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    
    # 查找B表的列索引
    price_col_b = find_column_index(ws_b, "销售价")
    qty_col_b = find_column_index(ws_b, "数量")
    amount_col_b = find_column_index(ws_b, "销售金额")
    
    if not price_col_b:
        error_msg = "错误：在B表中未找到'销售价'列"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    if not qty_col_b:
        error_msg = "错误：在B表中未找到'数量'列"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    if not amount_col_b:
        error_msg = "错误：在B表中未找到'销售金额'列"
        log(error_msg)
        if status_callback is None:
            raise ValueError(error_msg)
        return False
    
    log("列索引查找完成")
    log(f"A表：进价={price_col_a}, 进货数量={qty_col_a}, 进货金额={amount_col_a}")
    log(f"B表：销售价={price_col_b}, 数量={qty_col_b}, 销售金额={amount_col_b}")
    
    # 读取A表的所有数据行（从第2行开始）
    records_a = []
    for row_idx in range(2, ws_a.max_row + 1):
        price = get_cell_value(ws_a, row_idx, price_col_a)
        qty = get_cell_value(ws_a, row_idx, qty_col_a)
        amount = get_cell_value(ws_a, row_idx, amount_col_a)
        
        # 如果三个值都不为空，则记录（需要三个值都存在才能进行比较）
        if price is not None and qty is not None and amount is not None:
            # 转换为数值类型进行比较（如果可能）
            try:
                price = float(price) if price is not None else None
            except (ValueError, TypeError):
                pass
            try:
                qty = float(qty) if qty is not None else None
            except (ValueError, TypeError):
                pass
            try:
                amount = float(amount) if amount is not None else None
            except (ValueError, TypeError):
                pass
            
            records_a.append({
                'row': row_idx,
                'price': price,
                'qty': qty,
                'amount': amount
            })
    
    # 读取B表的所有数据行（从第2行开始）
    records_b = []
    for row_idx in range(2, ws_b.max_row + 1):
        price = get_cell_value(ws_b, row_idx, price_col_b)
        qty = get_cell_value(ws_b, row_idx, qty_col_b)
        amount = get_cell_value(ws_b, row_idx, amount_col_b)
        
        # 如果三个值都不为空，则记录（需要三个值都存在才能进行比较）
        if price is not None and qty is not None and amount is not None:
            # 转换为数值类型进行比较（如果可能）
            try:
                price = float(price) if price is not None else None
            except (ValueError, TypeError):
                pass
            try:
                qty = float(qty) if qty is not None else None
            except (ValueError, TypeError):
                pass
            try:
                amount = float(amount) if amount is not None else None
            except (ValueError, TypeError):
                pass
            
            records_b.append({
                'row': row_idx,
                'price': price,
                'qty': qty,
                'amount': amount
            })
    
    log(f"A表共有 {len(records_a)} 条数据记录")
    log(f"B表共有 {len(records_b)} 条数据记录")
    
    # 使用multiset匹配：统计每个三元组出现的次数
    # 创建键：使用元组(price, qty, amount)作为键
    # 对于浮点数，使用四舍五入到2位小数来避免精度问题
    def normalize_value(value):
        """规范化数值，用于比较"""
        if value is None:
            return None
        if isinstance(value, (int, float)):
            # 四舍五入到2位小数
            return round(float(value), 2)
        return value
    
    def make_key(record):
        return (
            normalize_value(record['price']),
            normalize_value(record['qty']),
            normalize_value(record['amount'])
        )
    
    counter_a = Counter()
    counter_b = Counter()
    
    for record in records_a:
        key = make_key(record)
        counter_a[key] += 1
    
    for record in records_b:
        key = make_key(record)
        counter_b[key] += 1
    
    # 找出未匹配的记录
    # 对于每个键，计算A表和B表中出现的次数差
    unmatched_rows_a = []
    unmatched_rows_b = []
    
    # 所有出现过的键
    all_keys = set(counter_a.keys()) | set(counter_b.keys())
    
    for key in all_keys:
        count_a = counter_a.get(key, 0)
        count_b = counter_b.get(key, 0)
        
        if count_a > count_b:
            # A表中多出的记录需要标记
            # 找到这些记录并标记
            remaining = count_a - count_b
            for record in records_a:
                if make_key(record) == key and remaining > 0:
                    unmatched_rows_a.append(record['row'])
                    remaining -= 1
        
        if count_b > count_a:
            # B表中多出的记录需要标记
            remaining = count_b - count_a
            for record in records_b:
                if make_key(record) == key and remaining > 0:
                    unmatched_rows_b.append(record['row'])
                    remaining -= 1
    
    log(f"\n匹配结果：")
    log(f"A表中未匹配的记录数：{len(unmatched_rows_a)}")
    log(f"B表中未匹配的记录数：{len(unmatched_rows_b)}")
    
    if unmatched_rows_a:
        log(f"A表未匹配的行号：{sorted(unmatched_rows_a)}")
        # 显示未匹配记录的详细信息
        log("A表未匹配记录的详细信息：")
        for row_idx in sorted(unmatched_rows_a):
            record = next(r for r in records_a if r['row'] == row_idx)
            log(f"  行{row_idx}: 进价={record['price']}, 进货数量={record['qty']}, 进货金额={record['amount']}")
    
    if unmatched_rows_b:
        log(f"B表未匹配的行号：{sorted(unmatched_rows_b)}")
        # 显示未匹配记录的详细信息
        log("B表未匹配记录的详细信息：")
        for row_idx in sorted(unmatched_rows_b):
            record = next(r for r in records_b if r['row'] == row_idx)
            log(f"  行{row_idx}: 销售价={record['price']}, 数量={record['qty']}, 销售金额={record['amount']}")
    
    # 标红未匹配的行
    red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    
    # 标红A表
    if unmatched_rows_a:
        max_col_a = ws_a.max_column if ws_a.max_column else 1
        # 如果max_column为None，至少标红到已找到的列
        max_col_a = max(max_col_a, price_col_a, qty_col_a, amount_col_a)
        for row_idx in unmatched_rows_a:
            for col_idx in range(1, max_col_a + 1):
                cell = ws_a.cell(row=row_idx, column=col_idx)
                cell.fill = red_fill
    
    # 标红B表
    if unmatched_rows_b:
        max_col_b = ws_b.max_column if ws_b.max_column else 1
        # 如果max_column为None，至少标红到已找到的列
        max_col_b = max(max_col_b, price_col_b, qty_col_b, amount_col_b)
        for row_idx in unmatched_rows_b:
            for col_idx in range(1, max_col_b + 1):
                cell = ws_b.cell(row=row_idx, column=col_idx)
                cell.fill = red_fill
    
    # 生成输出文件名
    file_a_name = Path(path_a).stem
    file_b_name = Path(path_b).stem
    file_a_ext = Path(path_a).suffix
    file_b_ext = Path(path_b).suffix
    
    output_file_a = os.path.join(save_path, f"{file_a_name}_找不同{file_a_ext}")
    output_file_b = os.path.join(save_path, f"{file_b_name}_找不同{file_b_ext}")
    
    # 保存文件
    log(f"\n正在保存A表到：{output_file_a}")
    wb_a.save(output_file_a)
    
    log(f"正在保存B表到：{output_file_b}")
    wb_b.save(output_file_b)
    
    log("\n处理完成！")
    return True


class CompareExcelGUI(tk.Tk):
    """比较Excel表格的GUI界面"""
    def __init__(self):
        super().__init__()
        self.title("比较Excel表格并标记差异")
        self.geometry("600x180")
        self.resizable(False, False)
        
        self.file1_path = tk.StringVar()
        self.file2_path = tk.StringVar()
        self.save_path = tk.StringVar()
        
        self._build_ui()
    
    def _build_ui(self):
        """构建UI界面"""
        pad = 6
        
        frame_main = ttk.Frame(self, padding=pad)
        frame_main.pack(fill="both", expand=True)
        
        # 文件1
        ttk.Label(frame_main, text="表格文件1:").grid(column=0, row=0, sticky="w", pady=2)
        ttk.Entry(frame_main, textvariable=self.file1_path, width=50).grid(column=1, row=0, padx=4, pady=2)
        ttk.Button(frame_main, text="选择文件", command=self.choose_file1).grid(column=2, row=0, pady=2)
        
        # 文件2
        ttk.Label(frame_main, text="表格文件2:").grid(column=0, row=1, sticky="w", pady=2)
        ttk.Entry(frame_main, textvariable=self.file2_path, width=50).grid(column=1, row=1, padx=4, pady=2)
        ttk.Button(frame_main, text="选择文件", command=self.choose_file2).grid(column=2, row=1, pady=2)
        
        # 保存路径
        ttk.Label(frame_main, text="保存路径:").grid(column=0, row=2, sticky="w", pady=2)
        ttk.Entry(frame_main, textvariable=self.save_path, width=50).grid(column=1, row=2, padx=4, pady=2)
        ttk.Button(frame_main, text="选择目录", command=self.choose_save_path).grid(column=2, row=2, pady=2)
        
        # 执行按钮
        frame_buttons = ttk.Frame(frame_main)
        frame_buttons.grid(column=0, columnspan=3, row=3, pady=8)
        ttk.Button(frame_buttons, text="开始比较", command=self.start_compare, width=17).pack(side="left", padx=4)
        ttk.Button(frame_buttons, text="退出", command=self.quit, width=17).pack(side="left", padx=4)
        
        # 状态栏
        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(frame_main, textvariable=self.status_var, foreground="gray").grid(column=0, columnspan=3, row=4, pady=3)
    
    def choose_file1(self):
        """选择第一个文件"""
        filename = filedialog.askopenfilename(
            title="选择第一个Excel表格文件",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if filename:
            self.file1_path.set(filename)
    
    def choose_file2(self):
        """选择第二个文件"""
        filename = filedialog.askopenfilename(
            title="选择第二个Excel表格文件",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if filename:
            self.file2_path.set(filename)
    
    def choose_save_path(self):
        """选择保存路径"""
        dirname = filedialog.askdirectory(title="选择保存目录")
        if dirname:
            self.save_path.set(dirname)
    
    def update_status(self, message):
        """更新状态栏"""
        self.status_var.set(message)
        self.update()
    
    def start_compare(self):
        """开始比较"""
        file1 = self.file1_path.get()
        file2 = self.file2_path.get()
        save_dir = self.save_path.get()
        
        if not file1 or not os.path.exists(file1):
            messagebox.showerror("错误", "请选择有效的第一个表格文件！")
            return
        
        if not file2 or not os.path.exists(file2):
            messagebox.showerror("错误", "请选择有效的第二个表格文件！")
            return
        
        if not save_dir:
            messagebox.showerror("错误", "请选择保存路径！")
            return
        
        try:
            self.status_var.set("正在处理...")
            self.update()
            
            # 调用比较函数
            success = compare_and_mark(
                file1, 
                file2, 
                save_dir,
                status_callback=self.update_status
            )
            
            if success:
                messagebox.showinfo("完成", "比较完成！未匹配的记录已标记为红色。")
                self.status_var.set("完成")
            else:
                self.status_var.set("处理失败")
        
        except Exception as e:
            traceback.print_exc()
            messagebox.showerror("错误", f"处理过程中发生错误：\n{e}")
            self.status_var.set(f"错误：{str(e)}")


def main():
    """主函数：支持GUI和命令行两种模式"""
    if len(sys.argv) == 4:
        # 命令行模式
        path_file1 = sys.argv[1]
        path_file2 = sys.argv[2]
        save_path = sys.argv[3]
        compare_and_mark(path_file1, path_file2, save_path)
    else:
        # GUI模式
        app = CompareExcelGUI()
        app.mainloop()


if __name__ == "__main__":
    main()
